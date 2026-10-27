# -*- coding: utf-8 -*-
"""My PDF Studio — Windows PDF editor. Offline after installation."""
import os
import sys
import tempfile
import tkinter as tk
from tkinter import ttk, filedialog, messagebox, simpledialog
from pathlib import Path

import fitz
from PIL import Image, ImageTk
from docx import Document
from docx.shared import Inches, Pt

APP = "My PDF Studio"
BG = "#eef2f7"
WHITE = "#ffffff"
NAVY = "#172b49"
BLUE = "#2866dc"
FONTFILE = Path(os.environ.get("WINDIR", "C:/Windows")) / "Fonts" / "arial.ttf"

def pdf_font(page):
    if FONTFILE.is_file():
        name = "mypdfarial"
        try:
            page.insert_font(fontname=name, fontfile=str(FONTFILE))
            return name
        except Exception:
            pass
    return "helv"

def draw_text(page, x, y, value, size=12):
    return page.insert_text((x, y), value, fontsize=size, fontname=pdf_font(page), color=(0,0,0), overlay=True)

class Studio:
    def __init__(self, root):
        self.root=root
        root.title(APP)
        root.geometry("1270x820")
        root.minsize(900,580)
        root.configure(bg=BG)
        self.doc=None
        self.path=None
        self.index=0
        self.scale=1.2
        self.changed=False
        self.backups=[]
        self.selected=None
        self.image=None
        self.line_data=[]
        self.create_ui()
        self.root.protocol("WM_DELETE_WINDOW",self.close)
        self.root.bind("<Control-o>",lambda e:self.open())
        self.root.bind("<Control-s>",lambda e:self.save())
        self.root.bind("<Control-p>",lambda e:self.print_file())
        self.root.bind("<Control-z>",lambda e:self.undo())
        self.root.bind("<Prior>",lambda e:self.prev())
        self.root.bind("<Next>",lambda e:self.next())
        self.status("Откройте PDF для просмотра или редактирования")
    def ui_button(self,parent,title,fn):
        return tk.Button(parent,text=title,command=fn,bg=WHITE,fg=NAVY,activebackground="#dce8ff",
                         font=("Segoe UI",10),relief="flat",padx=10,pady=8,cursor="hand2")
    def create_ui(self):
        top=tk.Frame(self.root,bg=NAVY)
        top.pack(fill="x")
        tk.Label(top,text="▣   My PDF Studio",font=("Segoe UI",17,"bold"),bg=NAVY,fg=WHITE).pack(side="left",padx=18,pady=14)
        tk.Label(top,text="PDF · DOCX · Печать",bg=NAVY,fg="#a4c7ff",font=("Segoe UI",10)).pack(side="right",padx=20)
        bar=tk.Frame(self.root,bg="#e3eaf3",padx=8,pady=6)
        bar.pack(fill="x")
        actions=[("Открыть PDF",self.open),("Новый",self.new),("Сохранить PDF",self.save),
                 ("PDF как...",lambda:self.save(True)),("В Word DOCX",self.export_word),
                 ("Печать",self.print_file),("Отменить",self.undo)]
        for text,fn in actions: self.ui_button(bar,text,fn).pack(side="left",padx=2)
        base=tk.PanedWindow(self.root,orient="horizontal",sashwidth=5,bg="#d8e0ec",bd=0)
        base.pack(fill="both",expand=True,padx=12,pady=10)
        left=tk.Frame(base,bg=WHITE,width=190)
        base.add(left,minsize=140)
        tk.Label(left,text="СТРАНИЦЫ",fg=NAVY,bg=WHITE,font=("Segoe UI",10,"bold")).pack(anchor="w",padx=12,pady=10)
        self.page_list=tk.Listbox(left,font=("Segoe UI",11),borderwidth=0,highlightthickness=0,
                                  selectbackground=BLUE,selectforeground=WHITE)
        self.page_list.pack(fill="both",expand=True,padx=9,pady=(0,10))
        self.page_list.bind("<<ListboxSelect>>",self.pick_page)
        center=tk.Frame(base,bg=BG)
        base.add(center,minsize=400)
        navigation=tk.Frame(center,bg=BG)
        navigation.pack(fill="x",pady=(2,9))
        for label,fn in [("◀",self.prev),("▶",self.next),("−",self.zoom_out),("+",self.zoom_in)]:
            self.ui_button(navigation,label,fn).pack(side="left",padx=2)
        self.nav_label=tk.Label(navigation,text="0 / 0",bg=BG,fg=NAVY,font=("Segoe UI",11,"bold"))
        self.nav_label.pack(side="left",padx=12)
        tk.Label(navigation,text="Щелчок по строке — редактирование",bg=BG,fg="#68778e").pack(side="right")
        self.canvas=tk.Canvas(center,bg="#a9b6c6",highlightthickness=0)
        self.canvas.pack(side="left",fill="both",expand=True)
        sc=tk.Scrollbar(center,orient="vertical",command=self.canvas.yview)
        sc.pack(side="right",fill="y")
        self.canvas.configure(yscrollcommand=sc.set)
        self.canvas.bind("<Button-1>",self.canvas_click)
        self.canvas.bind("<MouseWheel>",lambda e:self.canvas.yview_scroll(int(-1*(e.delta/120)),"units"))
        self.canvas.bind("<Configure>",lambda e:self.redraw())
        right=tk.Frame(base,bg=WHITE,width=280)
        base.add(right,minsize=230)
        tk.Label(right,text="РЕДАКТИРОВАНИЕ",fg=NAVY,bg=WHITE,font=("Segoe UI",11,"bold")).pack(anchor="w",padx=14,pady=(16,4))
        tk.Label(right,text="Выберите строку на странице\nили добавьте новую надпись.",fg="#65758d",
                 bg=WHITE,justify="left",font=("Segoe UI",10)).pack(anchor="w",padx=14,pady=(0,12))
        self.selection=tk.Label(right,text="Строка не выбрана",fg=NAVY,bg="#eef3fa",wraplength=244,justify="left",
                                padx=12,pady=10)
        self.selection.pack(fill="x",padx=12)
        self.edit=tk.Text(right,wrap="word",font=("Segoe UI",12),height=10,bg="#f8fafd",bd=0,highlightthickness=1,
                          highlightbackground="#d9e4f1",padx=9,pady=10)
        self.edit.pack(fill="both",expand=True,padx=12,pady=13)
        self.ui_button(right,"Заменить выбранную строку",self.replace).pack(fill="x",padx=12,pady=4)
        self.ui_button(right,"Удалить выбранную строку",self.delete).pack(fill="x",padx=12,pady=4)
        self.ui_button(right,"Добавить текст по щелчку",self.add_text).pack(fill="x",padx=12,pady=4)
        self.ui_button(right,"Повернуть страницу",self.rotate).pack(fill="x",padx=12,pady=4)
        self.ui_button(right,"Добавить страницу",self.add_page).pack(fill="x",padx=12,pady=4)
        self.ui_button(right,"Удалить страницу",self.remove_page).pack(fill="x",padx=12,pady=(4,16))
        self.statusbar=tk.Label(self.root,bg="#dfe8f5",fg=NAVY,anchor="w",font=("Segoe UI",9),padx=12,pady=7)
        self.statusbar.pack(fill="x")
    def status(self,msg):
        self.statusbar.configure(text=msg)
    def ensure(self):
        if not self.doc:
            messagebox.showinfo(APP,"Сначала откройте PDF или создайте новый документ.")
            return False
        return True
    def dirty_guard(self):
        if not self.changed: return True
        ans=messagebox.askyesnocancel(APP,"Сохранить изменения перед закрытием документа?")
        if ans is None: return False
        if ans: return self.save()
        return True
    def open(self):
        if not self.dirty_guard(): return
        path=filedialog.askopenfilename(title="Открыть PDF",filetypes=[("PDF","*.pdf"),("Все файлы","*.*")])
        if not path:return
        try:
            doc=fitz.open(path)
            if doc.needs_pass:
                password=simpledialog.askstring(APP,"Введите пароль PDF",show="*")
                if not password or not doc.authenticate(password):
                    doc.close();messagebox.showerror(APP,"Пароль неверный.");return
            if self.doc:self.doc.close()
            self.doc=doc;self.path=path;self.index=0;self.changed=False;self.backups=[]
            self.refresh()
        except Exception as e:messagebox.showerror("Ошибка открытия",str(e))
    def new(self):
        if not self.dirty_guard():return
        if self.doc:self.doc.close()
        self.doc=fitz.open();self.doc.new_page(width=595,height=842)
        self.path=None;self.index=0;self.backups=[];self.changed=True;self.refresh()
    def refresh(self):
        self.selected=None
        self.selection.configure(text="Строка не выбрана")
        self.edit.delete("1.0","end")
        self.page_list.delete(0,"end")
        if self.doc:
            for n in range(len(self.doc)):self.page_list.insert("end",f"  {n+1}")
            self.page_list.selection_set(self.index)
        self.redraw()
    def redraw(self):
        if not self.doc or len(self.doc)==0: return
        try:
            page=self.doc[self.index]
            pix=page.get_pixmap(matrix=fitz.Matrix(self.scale,self.scale),alpha=False)
            pic=Image.frombytes("RGB",(pix.width,pix.height),pix.samples)
            self.image=ImageTk.PhotoImage(pic)
            self.canvas.delete("all")
            width=max(self.canvas.winfo_width(),pix.width+32)
            x=(width-pix.width)/2
            self.canvas.create_image(x,16,image=self.image,anchor="nw")
            self.xoffset=x;self.yoffset=16
            self.canvas.configure(scrollregion=(0,0,width,pix.height+32))
            self.line_data=[]
            words=page.get_text("dict")
            for block in words["blocks"]:
                if "lines" not in block: continue
                for line in block["lines"]:
                    spans=line.get("spans",[])
                    text="".join(s.get("text","") for s in spans).strip()
                    if text:
                        self.line_data.append((text,fitz.Rect(line["bbox"]),spans[0].get("size",12)))
            self.nav_label.configure(text=f"{self.index+1} / {len(self.doc)}")
            self.root.title(APP+(" — "+Path(self.path).name if self.path else " — Новый документ")+("*" if self.changed else ""))
        except Exception as e:self.status("Ошибка отображения: "+str(e))
    def pick_page(self,event=None):
        if not self.doc:return
        selection=self.page_list.curselection()
        if selection:
            idx=selection[0]
            if idx!=self.index:
                self.index=idx;self.selected=None;self.redraw()
    def navigate(self,n):
        if not self.ensure():return
        self.index=max(0,min(len(self.doc)-1,self.index+n))
        self.page_list.selection_clear(0,"end")
        self.page_list.selection_set(self.index)
        self.page_list.see(self.index)
        self.selected=None;self.redraw()
    def prev(self):self.navigate(-1)
    def next(self):self.navigate(1)
    def zoom_in(self):self.scale=min(3.5,self.scale+0.15);self.redraw()
    def zoom_out(self):self.scale=max(0.5,self.scale-0.15);self.redraw()
    def canvas_click(self,event):
        if not self.doc:return
        px=(self.canvas.canvasx(event.x)-self.xoffset)/self.scale
        py=(self.canvas.canvasy(event.y)-self.yoffset)/self.scale
        for text,rect,size in self.line_data:
            if rect + (-2,-2,2,2) if False else False: pass
            if (rect.x0-3<=px<=rect.x1+3 and rect.y0-3<=py<=rect.y1+3):
                self.selected=(text,rect,size)
                self.selection.configure(text=f"Выбрано: {text[:120]}")
                self.edit.delete("1.0","end")
                self.edit.insert("1.0",text)
                self.status("Строка выбрана: можно заменить или удалить")
                return
        self.selected=None
        self.last_click=(px,py)
        self.selection.configure(text=f"Точка: {px:.0f}, {py:.0f} — добавление текста")
    def snapshot(self):
        self.backups.append(self.doc.tobytes(garbage=3,deflate=True))
        if len(self.backups)>8:self.backups.pop(0)
    def undo(self):
        if not self.ensure() or not self.backups:return
        try:
            data=self.backups.pop()
            self.doc.close();self.doc=fitz.open(stream=data,filetype="pdf")
            self.changed=True;self.refresh();self.status("Последнее действие отменено")
        except Exception as e:messagebox.showerror(APP,str(e))
    def erase_selected(self):
        text,rect,size=self.selected
        page=self.doc[self.index]
        page.add_redact_annot(rect,fill=(1,1,1))
        page.apply_redactions(images=0,graphics=0)
        return rect,size
    def replace(self):
        if not self.ensure() or not self.selected:
            messagebox.showinfo(APP,"Щёлкните по строке текста на странице.")
            return
        value=self.edit.get("1.0","end-1c")
        try:
            self.snapshot()
            rect,size=self.erase_selected()
            if value.strip():draw_text(self.doc[self.index],rect.x0,rect.y1-2,value,min(size,28))
            self.changed=True;self.refresh();self.status("Текст заменён. Проверьте расположение перед сохранением.")
        except Exception as e:messagebox.showerror(APP,str(e))
    def delete(self):
        if not self.ensure() or not self.selected:
            messagebox.showinfo(APP,"Щёлкните по строке для удаления.")
            return
        try:
            self.snapshot();self.erase_selected()
            self.changed=True;self.refresh();self.status("Строка удалена")
        except Exception as e:messagebox.showerror(APP,str(e))
    def add_text(self):
        if not self.ensure():return
        value=self.edit.get("1.0","end-1c").strip()
        if not value:
            messagebox.showinfo(APP,"Введите новый текст в поле справа.");return
        xy=getattr(self,"last_click",(55,80))
        size=simpledialog.askinteger(APP,"Размер шрифта",initialvalue=12,minvalue=7,maxvalue=50)
        if size is None:return
        try:
            self.snapshot()
            for i,line in enumerate(value.splitlines()):
                draw_text(self.doc[self.index],xy[0],xy[1]+size+i*(size*1.35),line,size)
            self.changed=True;self.refresh();self.status("Новый текст добавлен")
        except Exception as e:messagebox.showerror(APP,str(e))
    def rotate(self):
        if not self.ensure():return
        self.snapshot();p=self.doc[self.index];p.set_rotation((p.rotation+90)%360)
        self.changed=True;self.redraw()
    def add_page(self):
        if not self.ensure():return
        self.snapshot()
        self.doc.new_page(width=595,height=842,pno=self.index+1)
        self.index+=1;self.changed=True;self.refresh()
    def remove_page(self):
        if not self.ensure():return
        if len(self.doc)<=1:
            messagebox.showinfo(APP,"Последнюю страницу удалить нельзя.");return
        self.snapshot();self.doc.delete_page(self.index)
        self.index=min(self.index,len(self.doc)-1);self.changed=True;self.refresh()
    def save(self,as_new=False):
        if not self.ensure():return False
        path=self.path if self.path and not as_new else filedialog.asksaveasfilename(
            defaultextension=".pdf",filetypes=[("PDF","*.pdf")],initialfile="Документ.pdf")
        if not path:return False
        if not path.lower().endswith(".pdf"):path+=".pdf"
        try:
            temp=Path(path).with_suffix(".mypdf-temp.pdf")
            self.doc.save(str(temp),garbage=3,deflate=True)
            os.replace(temp,path)
            self.path=path;self.changed=False;self.redraw()
            self.status("PDF сохранён: "+path)
            return True
        except Exception as e:
            messagebox.showerror("Ошибка сохранения",str(e))
            return False
    def export_word(self):
        if not self.ensure():return
        target=filedialog.asksaveasfilename(title="Экспорт Word",defaultextension=".docx",
                 filetypes=[("Документ Word","*.docx")],initialfile="Документ.docx")
        if not target:return
        try:
            doc=Document()
            section=doc.sections[0]
            section.top_margin=Inches(0.75);section.bottom_margin=Inches(0.75)
            for k,page in enumerate(self.doc):
                if k:doc.add_page_break()
                blocks=page.get_text("blocks")
                if not any(str(b[4]).strip() for b in blocks if len(b)>4 and b[6]==0):
                    snap=page.get_pixmap(matrix=fitz.Matrix(1.2,1.2),alpha=False)
                    from io import BytesIO
                    buff=BytesIO(snap.tobytes("png"))
                    doc.add_picture(buff,width=Inches(6.5))
                else:
                    for raw in blocks:
                        if len(raw)>6 and raw[6]==0:
                            for text in str(raw[4]).splitlines():
                                if text.strip():doc.add_paragraph(text)
            doc.save(target)
            self.status("Word DOCX сохранён: "+target)
            messagebox.showinfo(APP,"Документ сохранён в DOCX. Сложная PDF-вёрстка может отличаться.")
        except Exception as e:messagebox.showerror("Ошибка экспорта Word",str(e))
    def print_file(self):
        if not self.ensure():return
        try:
            if self.changed or not self.path:
                tmp=Path(tempfile.gettempdir())/"MyPDFStudio_Print.pdf"
                self.doc.save(str(tmp),garbage=3,deflate=True)
                path=str(tmp)
            else:path=self.path
            os.startfile(path,"print")
            self.status("Отправлено на печать через приложение PDF по умолчанию")
        except Exception as e:
            messagebox.showerror(APP,"Не удалось отправить на печать. Проверьте, что в Windows задано приложение для PDF.\n"+str(e))
    def close(self):
        if not self.dirty_guard():return
        if self.doc:self.doc.close()
        self.root.destroy()

if __name__=="__main__":
    root=tk.Tk()
    Studio(root)
    root.mainloop()
